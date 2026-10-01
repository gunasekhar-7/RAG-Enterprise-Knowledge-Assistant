import io

import docx
import fitz
import pytest

from app.ingestion.exceptions import CorruptDocumentError
from app.ingestion.models import DocumentMetadata, DocumentType
from app.ingestion.parsers.base import BaseParser
from app.ingestion.parsers.docx import DocxParser
from app.ingestion.parsers.pdf import PDFParser
from app.ingestion.parsers.text import TextParser


def create_sample_pdf_bytes() -> bytes:
    """Helper to generate an in-memory multi-page PDF."""
    doc = fitz.open()
    doc.set_metadata({"title": "Employee Handbook", "author": "HR Department"})

    page1 = doc.new_page()
    page1.insert_text((50, 50), "Welcome to the company handbook. Page 1 content.")

    page2 = doc.new_page()
    page2.insert_text((50, 50), "Health and security policies. Page 2 content.")

    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes


def create_sample_docx_bytes() -> bytes:
    """Helper to generate an in-memory DOCX document."""
    doc = docx.Document()
    doc.core_properties.title = "Engineering SOP"
    doc.core_properties.author = "Platform Team"

    doc.add_heading("Platform Operations SOP", level=1)
    doc.add_paragraph("All deployments must follow the canary procedure.")

    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Component"
    table.rows[0].cells[1].text = "Version"
    table.rows[1].cells[0].text = "FastAPI"
    table.rows[1].cells[1].text = "0.115.0"

    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()


def test_clean_text() -> None:
    raw = "Line 1 \r\n\r\n\r\n\r\nLine 2   \x00\x07  \nLine 3"
    cleaned = BaseParser.clean_text(raw)
    assert "\r" not in cleaned
    assert "\x00" not in cleaned
    assert "\n\n\n" not in cleaned  # collapsed to double newline
    assert cleaned == "Line 1\n\nLine 2\nLine 3"


def test_pdf_parser_page_preservation() -> None:
    pdf_bytes = create_sample_pdf_bytes()
    parser = PDFParser()
    meta = DocumentMetadata(
        filename="handbook.pdf",
        sanitized_filename="handbook.pdf",
        file_type=DocumentType.PDF,
        mime_type="application/pdf",
        file_size_bytes=len(pdf_bytes),
        sha256_hash="a" * 64,
    )

    pages, extracted_meta = parser.parse(pdf_bytes, meta)

    # Validate exact page numbers preserved for citations
    assert len(pages) == 2
    assert pages[0].page_number == 1
    assert "Welcome to the company handbook" in pages[0].text
    assert pages[1].page_number == 2
    assert "Health and security policies" in pages[1].text

    # Validate embedded metadata
    assert extracted_meta["title"] == "Employee Handbook"
    assert extracted_meta["author"] == "HR Department"


def test_pdf_parser_corrupt_file() -> None:
    parser = PDFParser()
    meta = DocumentMetadata(
        filename="corrupt.pdf",
        sanitized_filename="corrupt.pdf",
        file_type=DocumentType.PDF,
        mime_type="application/pdf",
        file_size_bytes=20,
        sha256_hash="b" * 64,
    )

    with pytest.raises(CorruptDocumentError):
        parser.parse(b"%PDF-not_a_valid_pdf_stream", meta)


def test_docx_parser() -> None:
    docx_bytes = create_sample_docx_bytes()
    parser = DocxParser()
    meta = DocumentMetadata(
        filename="sop.docx",
        sanitized_filename="sop.docx",
        file_type=DocumentType.DOCX,
        mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        file_size_bytes=len(docx_bytes),
        sha256_hash="c" * 64,
    )

    pages, extracted_meta = parser.parse(docx_bytes, meta)

    assert len(pages) >= 1
    assert pages[0].page_number == 1
    combined_text = "\n".join(p.text for p in pages)
    assert "Platform Operations SOP" in combined_text
    assert "FastAPI | 0.115.0" in combined_text
    assert extracted_meta["title"] == "Engineering SOP"
    assert extracted_meta["author"] == "Platform Team"


def test_text_parser_markdown() -> None:
    md_content = b"""# Information Security Policy

All production database access must go through AWS IAM roles.

## Password Requirements
Passwords must have at least 16 characters.
"""

    parser = TextParser()
    meta = DocumentMetadata(
        filename="security.md",
        sanitized_filename="security.md",
        file_type=DocumentType.MARKDOWN,
        mime_type="text/markdown",
        file_size_bytes=len(md_content),
        sha256_hash="d" * 64,
    )

    pages, extracted_meta = parser.parse(md_content, meta)

    assert len(pages) == 1
    assert pages[0].page_number == 1
    assert "Information Security Policy" in pages[0].text
    assert extracted_meta["title"] == "Information Security Policy"
